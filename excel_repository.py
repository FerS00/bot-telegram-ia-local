"""Acceso de solo lectura al libro ShopIA_Base_Datos_RAG.xlsx.

El libro es la fuente autoritativa de precios, stock, envíos, promociones,
políticas y documentos RAG. Nunca se escribe. Las columnas se resuelven por
nombre de cabecera, así que el orden de columnas del libro no importa.
"""

from __future__ import annotations

import re
import unicodedata
import zipfile
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

COMPRABLE_ESTADO = "activo"
CENTAVO = Decimal("0.01")

# Hoja -> columnas que el código necesita. Las demás columnas se ignoran.
REQUIRED_SHEETS: dict[str, tuple[str, ...]] = {
    "Productos": (
        "producto_id", "sku", "categoria_id", "categoria", "marca", "modelo", "nombre",
        "descripcion_corta", "descripcion_larga", "precio_pen", "moneda",
        "uso_recomendado", "garantia_meses", "estado", "etiquetas",
    ),
    "Inventario": ("producto_id", "almacen_id", "stock_disponible"),
    "Especificaciones": ("producto_id", "atributo", "valor", "unidad"),
    "Documentos_RAG": (
        "documento_id", "tipo_documento", "hoja_origen", "registro_id", "titulo",
        "contenido", "palabras_clave", "prioridad",
    ),
    "Categorias": ("categoria_id", "nombre", "descripcion"),
    "Envios": (
        "zona_id", "departamento", "provincia", "distrito_o_ciudad", "tarifa_base_pen",
        "plazo_min_dias_habiles", "plazo_max_dias_habiles", "monto_envio_gratis_desde_pen",
    ),
    "Promociones": (
        "codigo", "nombre", "tipo_descuento", "porcentaje", "monto_fijo_pen",
        "monto_minimo_pen", "descuento_maximo_pen", "categoria_id", "fecha_inicio",
        "fecha_fin", "estado",
    ),
    "Politicas": ("politica_id", "tema", "descripcion"),
    "FAQ": ("faq_id", "pregunta", "respuesta"),
    "Guias_Compra": ("guia_id", "titulo", "contenido"),
    "Ordenes_demo": (
        "pedido_id", "fecha_simulada", "estado", "subtotal_pen", "descuento_pen",
        "envio_pen", "total_pen", "tipo",
    ),
    "Detalle_Orden": ("pedido_id", "sku", "cantidad", "precio_unitario_pen", "subtotal_pen"),
}


class WorkbookError(RuntimeError):
    """El libro no existe, no se abre o le faltan hojas/columnas."""


@dataclass(frozen=True)
class Especificacion:
    atributo: str
    valor: str
    unidad: str | None


@dataclass(frozen=True)
class Producto:
    producto_id: str
    sku: str
    categoria_id: str
    categoria: str
    marca: str
    modelo: str
    nombre: str
    descripcion_corta: str
    descripcion_larga: str
    precio_pen: Decimal
    uso_recomendado: str
    etiquetas: str
    garantia_meses: int
    estado: str
    stock: int
    especificaciones: tuple[Especificacion, ...]

    @property
    def comprable(self) -> bool:
        return self.estado.strip().lower() == COMPRABLE_ESTADO and self.stock > 0


@dataclass(frozen=True)
class ResultadoBusqueda:
    producto: Producto
    puntaje: float
    coincidencias_fuertes: int


@dataclass(frozen=True)
class Categoria:
    categoria_id: str
    nombre: str
    descripcion: str


@dataclass(frozen=True)
class ZonaEnvio:
    zona_id: str
    departamento: str
    provincia: str
    distrito: str
    tarifa_base: Decimal
    plazo_min_dias: int
    plazo_max_dias: int
    gratis_desde: Decimal | None


@dataclass(frozen=True)
class Promocion:
    codigo: str
    nombre: str
    tipo_descuento: str
    porcentaje: Decimal | None
    monto_fijo: Decimal | None
    monto_minimo: Decimal
    descuento_maximo: Decimal | None
    categoria_id: str
    fecha_inicio: date
    fecha_fin: date
    estado: str


@dataclass(frozen=True)
class Documento:
    documento_id: str
    tipo_documento: str
    hoja_origen: str
    registro_id: str
    titulo: str
    contenido: str
    palabras_clave: str
    prioridad: str


@dataclass(frozen=True)
class LineaOrden:
    sku: str
    cantidad: int
    precio_unitario: Decimal
    subtotal: Decimal


@dataclass(frozen=True)
class OrdenDemo:
    pedido_id: str
    fecha: date | None
    estado: str
    subtotal: Decimal
    descuento: Decimal
    envio: Decimal
    total: Decimal
    tipo: str
    lineas: tuple[LineaOrden, ...]


@dataclass(frozen=True)
class _Datos:
    productos: tuple[Producto, ...]
    productos_por_sku: dict[str, Producto]
    categorias: tuple[Categoria, ...]
    zonas: tuple[ZonaEnvio, ...]
    promociones: tuple[Promocion, ...]
    documentos: tuple[Documento, ...]
    ordenes: dict[str, OrdenDemo]
    conteos: dict[str, int]


# ---------------------------------------------------------------- utilidades

def normalizar(texto: str | None) -> str:
    """Minúsculas, sin tildes y con espacios colapsados."""
    if not texto:
        return ""
    sin_tildes = unicodedata.normalize("NFKD", str(texto))
    sin_tildes = "".join(c for c in sin_tildes if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", sin_tildes.lower()).strip()


def _texto(valor: Any) -> str:
    return "" if valor is None else str(valor).strip()


def _decimal(valor: Any, columna: str, hoja: str) -> Decimal | None:
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    try:
        return Decimal(str(valor).strip()).quantize(CENTAVO)
    except (InvalidOperation, ValueError) as exc:
        raise WorkbookError(f"Valor numérico no válido en {hoja}.{columna}.") from exc


def _entero(valor: Any, columna: str, hoja: str) -> int:
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return 0
    try:
        return int(Decimal(str(valor).strip()))
    except (InvalidOperation, ValueError) as exc:
        raise WorkbookError(f"Valor entero no válido en {hoja}.{columna}.") from exc


def _fecha(valor: Any) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    return None


# ---------------------------------------------------------------- lectura

def _leer_hojas(ruta: Path) -> dict[str, list[dict[str, Any]]]:
    if not ruta.is_file():
        raise WorkbookError(f"No se encontró el libro Excel en la ruta configurada: {ruta.name}.")
    try:
        libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    except (OSError, zipfile.BadZipFile, InvalidFileException, KeyError) as exc:
        raise WorkbookError("El archivo no es un libro .xlsx legible.") from exc

    try:
        hojas_presentes = set(libro.sheetnames)
        faltantes = [h for h in REQUIRED_SHEETS if h not in hojas_presentes]
        if faltantes:
            raise WorkbookError(f"Faltan hojas en el libro: {', '.join(faltantes)}.")

        resultado: dict[str, list[dict[str, Any]]] = {}
        for hoja, columnas in REQUIRED_SHEETS.items():
            resultado[hoja] = _leer_hoja(libro[hoja], hoja, columnas)
        return resultado
    finally:
        libro.close()


def _leer_hoja(ws: Any, hoja: str, columnas: Sequence[str]) -> list[dict[str, Any]]:
    filas = ws.iter_rows(values_only=True)
    cabecera = next(filas, None)
    if cabecera is None:
        raise WorkbookError(f"La hoja {hoja} está vacía; falta la fila de cabeceras.")

    indices: dict[str, int] = {}
    for i, nombre in enumerate(cabecera):
        if isinstance(nombre, str) and nombre.strip():
            indices.setdefault(nombre.strip(), i)

    faltantes = [c for c in columnas if c not in indices]
    if faltantes:
        raise WorkbookError(f"Faltan columnas en {hoja}: {', '.join(faltantes)}.")

    registros: list[dict[str, Any]] = []
    for fila in filas:
        if fila is None or all(celda is None for celda in fila):
            continue
        registros.append({c: fila[indices[c]] if indices[c] < len(fila) else None for c in columnas})
    return registros


# ---------------------------------------------------------------- construcción

def _construir(hojas: dict[str, list[dict[str, Any]]]) -> _Datos:
    stock_por_producto: dict[str, int] = {}
    for fila in hojas["Inventario"]:
        pid = _texto(fila["producto_id"])
        stock_por_producto[pid] = stock_por_producto.get(pid, 0) + _entero(
            fila["stock_disponible"], "stock_disponible", "Inventario"
        )

    specs_por_producto: dict[str, list[Especificacion]] = {}
    for fila in hojas["Especificaciones"]:
        specs_por_producto.setdefault(_texto(fila["producto_id"]), []).append(
            Especificacion(
                atributo=_texto(fila["atributo"]),
                valor=_texto(fila["valor"]),
                unidad=_texto(fila["unidad"]) or None,
            )
        )

    productos: list[Producto] = []
    for fila in hojas["Productos"]:
        pid = _texto(fila["producto_id"])
        precio = _decimal(fila["precio_pen"], "precio_pen", "Productos")
        if precio is None:
            raise WorkbookError("Hay productos sin precio_pen en la hoja Productos.")
        productos.append(
            Producto(
                producto_id=pid,
                sku=_texto(fila["sku"]),
                categoria_id=_texto(fila["categoria_id"]),
                categoria=_texto(fila["categoria"]),
                marca=_texto(fila["marca"]),
                modelo=_texto(fila["modelo"]),
                nombre=_texto(fila["nombre"]),
                descripcion_corta=_texto(fila["descripcion_corta"]),
                descripcion_larga=_texto(fila["descripcion_larga"]),
                precio_pen=precio,
                uso_recomendado=_texto(fila["uso_recomendado"]),
                etiquetas=_texto(fila["etiquetas"]),
                garantia_meses=_entero(fila["garantia_meses"], "garantia_meses", "Productos"),
                estado=_texto(fila["estado"]),
                stock=stock_por_producto.get(pid, 0),
                especificaciones=tuple(specs_por_producto.get(pid, ())),
            )
        )

    zonas = tuple(
        ZonaEnvio(
            zona_id=_texto(f["zona_id"]),
            departamento=_texto(f["departamento"]),
            provincia=_texto(f["provincia"]),
            distrito=_texto(f["distrito_o_ciudad"]),
            tarifa_base=_decimal(f["tarifa_base_pen"], "tarifa_base_pen", "Envios") or Decimal("0"),
            plazo_min_dias=_entero(f["plazo_min_dias_habiles"], "plazo_min_dias_habiles", "Envios"),
            plazo_max_dias=_entero(f["plazo_max_dias_habiles"], "plazo_max_dias_habiles", "Envios"),
            gratis_desde=_decimal(
                f["monto_envio_gratis_desde_pen"], "monto_envio_gratis_desde_pen", "Envios"
            ),
        )
        for f in hojas["Envios"]
    )

    promociones = tuple(
        Promocion(
            codigo=_texto(f["codigo"]).upper(),
            nombre=_texto(f["nombre"]),
            tipo_descuento=_texto(f["tipo_descuento"]),
            porcentaje=_decimal(f["porcentaje"], "porcentaje", "Promociones"),
            monto_fijo=_decimal(f["monto_fijo_pen"], "monto_fijo_pen", "Promociones"),
            monto_minimo=_decimal(f["monto_minimo_pen"], "monto_minimo_pen", "Promociones")
            or Decimal("0"),
            descuento_maximo=_decimal(f["descuento_maximo_pen"], "descuento_maximo_pen", "Promociones"),
            categoria_id=_texto(f["categoria_id"]).upper(),
            fecha_inicio=_fecha(f["fecha_inicio"]) or date.min,
            fecha_fin=_fecha(f["fecha_fin"]) or date.max,
            estado=_texto(f["estado"]),
        )
        for f in hojas["Promociones"]
    )

    documentos = tuple(
        Documento(
            documento_id=_texto(f["documento_id"]),
            tipo_documento=_texto(f["tipo_documento"]),
            hoja_origen=_texto(f["hoja_origen"]),
            registro_id=_texto(f["registro_id"]),
            titulo=_texto(f["titulo"]),
            contenido=_texto(f["contenido"]),
            palabras_clave=_texto(f["palabras_clave"]),
            prioridad=_texto(f["prioridad"]),
        )
        for f in hojas["Documentos_RAG"]
    )

    lineas_por_pedido: dict[str, list[LineaOrden]] = {}
    for f in hojas["Detalle_Orden"]:
        lineas_por_pedido.setdefault(_texto(f["pedido_id"]), []).append(
            LineaOrden(
                sku=_texto(f["sku"]),
                cantidad=_entero(f["cantidad"], "cantidad", "Detalle_Orden"),
                precio_unitario=_decimal(f["precio_unitario_pen"], "precio_unitario_pen", "Detalle_Orden")
                or Decimal("0"),
                subtotal=_decimal(f["subtotal_pen"], "subtotal_pen", "Detalle_Orden") or Decimal("0"),
            )
        )
    ordenes: dict[str, OrdenDemo] = {}
    for f in hojas["Ordenes_demo"]:
        pid = _texto(f["pedido_id"])
        ordenes[pid] = OrdenDemo(
            pedido_id=pid,
            fecha=_fecha(f["fecha_simulada"]),
            estado=_texto(f["estado"]),
            subtotal=_decimal(f["subtotal_pen"], "subtotal_pen", "Ordenes_demo") or Decimal("0"),
            descuento=_decimal(f["descuento_pen"], "descuento_pen", "Ordenes_demo") or Decimal("0"),
            envio=_decimal(f["envio_pen"], "envio_pen", "Ordenes_demo") or Decimal("0"),
            total=_decimal(f["total_pen"], "total_pen", "Ordenes_demo") or Decimal("0"),
            tipo=_texto(f["tipo"]),
            lineas=tuple(lineas_por_pedido.get(pid, ())),
        )

    categorias = tuple(
        Categoria(_texto(f["categoria_id"]), _texto(f["nombre"]), _texto(f["descripcion"]))
        for f in hojas["Categorias"]
    )

    return _Datos(
        productos=tuple(productos),
        productos_por_sku={p.sku.upper(): p for p in productos},
        categorias=categorias,
        zonas=zonas,
        promociones=promociones,
        documentos=documentos,
        ordenes=ordenes,
        conteos={hoja: len(filas) for hoja, filas in hojas.items()},
    )


# ---------------------------------------------------------------- repositorio

class ExcelRepository:
    """Carga el libro una vez y expone consultas en memoria."""

    def __init__(self, ruta: Path | str) -> None:
        self._ruta = Path(ruta)
        self._datos: _Datos | None = None

    @property
    def ruta(self) -> Path:
        return self._ruta

    def load(self) -> "ExcelRepository":
        if self._datos is None:
            self._datos = _construir(_leer_hojas(self._ruta))
        return self

    @property
    def _d(self) -> _Datos:
        if self._datos is None:
            self.load()
        assert self._datos is not None
        return self._datos

    def conteos(self) -> dict[str, int]:
        """Filas por hoja requerida (para diagnóstico, sin contenido)."""
        return dict(self._d.conteos)

    def productos(self) -> list[Producto]:
        return list(self._d.productos)

    def producto(self, sku: str) -> Producto | None:
        return self._d.productos_por_sku.get(_texto(sku).upper())

    def categorias(self) -> list[Categoria]:
        return list(self._d.categorias)

    def categoria(self, texto: str) -> Categoria | None:
        clave = normalizar(texto)
        for cat in self._d.categorias:
            if clave in (normalizar(cat.categoria_id), normalizar(cat.nombre)):
                return cat
        return None

    def zonas_envio(self) -> list[ZonaEnvio]:
        return list(self._d.zonas)

    def zonas_por_destino(self, texto: str) -> tuple[str, list[ZonaEnvio]]:
        """Busca por distrito, luego provincia y luego departamento.

        Devuelve el nivel que coincidió ("distrito", "provincia", "departamento")
        y las zonas encontradas. Una lista vacía significa destino no hallado.
        """
        partes = [normalizar(p) for p in re.split(r"[,;]", texto or "") if normalizar(p)]
        if not partes:
            return "", []
        niveles = (
            ("distrito", lambda z: normalizar(z.distrito)),
            ("provincia", lambda z: normalizar(z.provincia)),
            ("departamento", lambda z: normalizar(z.departamento)),
        )
        for nivel, clave in niveles:
            encontradas = [z for z in self._d.zonas if any(clave(z) == p for p in partes)]
            if encontradas:
                return nivel, encontradas
        return "", []

    def zona_en_texto(self, texto: str) -> ZonaEnvio | None:
        """Detecta un distrito citado dentro de una frase libre (el más largo gana)."""
        base = f" {normalizar(texto)} "
        candidatas = sorted(
            (z for z in self._d.zonas if normalizar(z.distrito)),
            key=lambda z: len(normalizar(z.distrito)),
            reverse=True,
        )
        for zona in candidatas:
            if f" {normalizar(zona.distrito)} " in base:
                return zona
        return None

    def promociones(self) -> list[Promocion]:
        return list(self._d.promociones)

    def promocion(self, codigo: str) -> Promocion | None:
        clave = _texto(codigo).upper()
        for promo in self._d.promociones:
            if promo.codigo == clave:
                return promo
        return None

    def documentos(self) -> list[Documento]:
        return list(self._d.documentos)

    def pedido_demo(self, pedido_id: str) -> OrdenDemo | None:
        return self._d.ordenes.get(_texto(pedido_id).upper())

    def buscar(
        self,
        texto: str = "",
        *,
        categoria: str | None = None,
        marca: str | None = None,
        uso: str | None = None,
        especificaciones: Iterable[str] = (),
        max_price: Decimal | None = None,
        limite: int = 4,
        solo_comprables: bool = True,
    ) -> list[ResultadoBusqueda]:
        """Búsqueda léxica con filtros. Ordena por relevancia y luego por precio."""
        terminos = _tokens(texto)
        if texto and texto.strip() and not terminos:
            # Solo palabras vacías ("quiero", "hay"): no es una búsqueda, no devuelve todo el catálogo.
            return []
        filtro_cat = normalizar(categoria) if categoria else ""
        filtro_marca = normalizar(marca) if marca else ""
        filtro_uso = normalizar(uso) if uso else ""
        filtros_spec = [_tokens(s) for s in especificaciones if _tokens(s)]

        resultados: list[ResultadoBusqueda] = []
        for prod in self._d.productos:
            if solo_comprables and not prod.comprable:
                continue
            if max_price is not None and prod.precio_pen > max_price:
                continue
            if filtro_cat and filtro_cat not in (
                normalizar(prod.categoria_id), normalizar(prod.categoria)
            ):
                continue
            if filtro_marca and filtro_marca != normalizar(prod.marca):
                continue
            if filtro_uso and filtro_uso not in normalizar(f"{prod.uso_recomendado} {prod.etiquetas}"):
                continue
            specs_tokens = [_tokens(f"{e.atributo} {e.valor}") for e in prod.especificaciones]
            if not all(any(f <= st for st in specs_tokens) for f in filtros_spec):
                continue

            puntaje, fuertes = _puntuar(prod, terminos, specs_tokens)
            # Descartar coincidencias solo en descripción: generan ruido en búsquedas por precio.
            if terminos and puntaje < 2:
                continue
            resultados.append(ResultadoBusqueda(prod, puntaje, fuertes))

        resultados.sort(key=lambda r: (-r.puntaje, r.producto.precio_pen, r.producto.sku))
        return resultados[: max(limite, 0)] if limite else resultados

    def atributos_numericos(self, producto: Producto, atributo: str) -> Decimal | None:
        """Primer número del atributo indicado (por nombre normalizado)."""
        clave = normalizar(atributo)
        for esp in producto.especificaciones:
            if normalizar(esp.atributo) == clave:
                encontrado = re.search(r"\d+(?:[.,]\d+)?", esp.valor)
                if encontrado:
                    return Decimal(encontrado.group(0).replace(",", "."))
        return None


# ---------------------------------------------------------------- ranking

_STOPWORDS = frozenset(
    "a al con de del del el en es la las lo los me mi o para por que quiero busco "
    "necesito tengo tienes tiene hay un una unos unas y ya se sin sobre "
    "hasta menos mas cuanto cuesta precio soles sol pen s dame muestra muestrame "
    "recomienda recomiendame ver algo alguno otro otra".split()
)


# Equivalencias genéricas de lenguaje común (no son datos del libro).
_SINONIMOS = {
    "portatil": "laptop", "notebook": "laptop",
    "celular": "smartphone", "movil": "smartphone", "telefono": "smartphone",
    "auricular": "audifono", "headset": "audifono",
    "raton": "mouse", "tableta": "tablet",
}


def _tokens(texto: str) -> set[str]:
    return {_SINONIMOS.get(_raiz(t), _raiz(t)) for t in re.findall(r"[a-z0-9][a-z0-9./-]*", normalizar(texto))
            if t not in _STOPWORDS}


def _raiz(token: str) -> str:
    """Raíz simple para plurales del español (portátiles -> portatil)."""
    if len(token) > 4 and token.endswith("es"):
        return token[:-2]
    if len(token) > 3 and token.endswith("s"):
        return token[:-1]
    return token


def _puntuar(prod: Producto, terminos: set[str], specs_tokens: list[set[str]]) -> tuple[float, int]:
    if not terminos:
        return 1.0, 0
    fuerte = _tokens(f"{prod.nombre} {prod.marca} {prod.modelo} {prod.categoria}")
    media = _tokens(f"{prod.uso_recomendado} {prod.etiquetas}")
    debil = _tokens(f"{prod.descripcion_corta} {prod.descripcion_larga}")
    specs: set[str] = set().union(*specs_tokens) if specs_tokens else set()

    puntaje = 0.0
    fuertes = 0
    for t in terminos:
        if t in fuerte:
            puntaje += 3.0
            fuertes += 1
        elif t in media or t in specs:
            puntaje += 2.0
        elif t in debil:
            puntaje += 1.0
    return puntaje, fuertes
